"""
Build a stakeholder email (HTML, inline styles so Outlook keeps them) from a canary report workbook
written by fetch_canary_report.py.

Usage:
    python build_email_html.py canary_report_2026-09-01_2026-09-30.xlsx
    python build_email_html.py canary_report_2026-09-01_2026-09-30.xlsx --output september_email.html

Reads the Summary and Info sheets only; the workbook itself is the attachment with the full detail.
"""

import argparse
import html
import re
from collections import defaultdict
from datetime import datetime

from openpyxl import load_workbook

from fetch_canary_report import REGIONS

# Column order and display names; regions with the full canary suite go in the first matrix.
SUITE_REGIONS = {"na1": "NA1", "de": "DE", "au": "AU", "ausov": "AUSOV1", "eusov": "EUSOV1",
                 "uksov": "UKSOV1", "fedramp": "FEDRAMP", "uk3": "UK3"}
ALL_REGIONS = {"na1": "NA1", "de": "DE", "au": "AU", "jp": "JP", "ca": "CA", "uk": "UK", "uae1": "AE",
               "jo": "JO", "ausov": "AUSOV1", "kr": "KR", "eusov": "EUSOV1", "za": "ZA", "uksov": "UKSOV1",
               "fedramp": "FEDRAMP", "na3": "NA3", "uk3": "UK3", "eu3": "EU3"}
DATA_SYNC = ["ga-dsync-can", "ga-dspwd-can"]
GREEN, AMBER = 99.5, 95.0

NAVY, BORDER, STRIPE, MUTED = "#1f3864", "#cfd8dc", "#f4f6fa", "#90a4ae"
FONT = "font-family:Calibri,Arial,sans-serif;"
TH = f"border:1px solid {BORDER};padding:6px 8px;background:{NAVY};color:#ffffff;font-weight:600;font-size:13px;"
TD = f"border:1px solid {BORDER};padding:6px 8px;font-size:13px;vertical-align:top;"


def parse_args():
    p = argparse.ArgumentParser(description="Stakeholder email HTML from a canary report workbook")
    p.add_argument("report", help="canary_report_<start>_<end>.xlsx from fetch_canary_report.py")
    p.add_argument("--output", help="Output .html path (default: <report>_email.html)")
    return p.parse_args()


def base_name(region, canary):
    for prefix in REGIONS[region]["prefixes"]:
        if canary.startswith(prefix):
            return canary[len(prefix):]
    return canary


def short_reason(reason):
    """Top failure reason without per-run noise: counts, URLs and long tails."""
    reason = re.sub(r"\s*\(\d+x\)\s*$", "", reason or "")
    reason = re.split(r" \(at https?://| URL: |\. Error: | Call log:", reason)[0]
    return reason if len(reason) <= 160 else reason[:157] + "..."


def load_report(path):
    wb = load_workbook(path, read_only=True)
    info = [r[0] for r in wb["Info"].iter_rows(values_only=True) if r and r[0]]
    period = re.search(r"Date range \((.+?), inclusive\): (\d{4}-\d{2}-\d{2}) to (\d{4}-\d{2}-\d{2})", info[0])
    tz_name = period.group(1)
    start, end = (datetime.strptime(d, "%Y-%m-%d") for d in period.groups()[1:])
    rows = list(wb["Summary"].iter_rows(values_only=True))
    header = rows[0]
    cells = {}
    for values in rows[1:]:
        r = dict(zip(header, values))
        if r["Region"] not in ALL_REGIONS:
            continue
        cells[(base_name(r["Region"], r["Canary"]), r["Region"])] = {
            "canary": r["Canary"], "state": r["State"], "runs": r["Runs"] or 0, "passed": r["Passed"] or 0,
            "failed": r["Failed"] or 0, "pct": r["Pass %"], "reason": short_reason(r["Top failure reason"]),
        }
    return start, end, tz_name, cells


def colour(pct):
    return "#2e7d32" if pct >= GREEN else "#b7791f" if pct >= AMBER else "#c00000"


def pct_cell(cell):
    if not cell or cell["state"] == "NOT FOUND":
        return f'<td style="{TD}text-align:center;color:{MUTED};">&ndash;</td>'
    if cell["pct"] is None or not cell["runs"]:
        return f'<td style="{TD}text-align:center;color:{MUTED};font-size:11px;">stopped</td>'
    value = float(cell["pct"])
    return (f'<td style="{TD}text-align:center;font-weight:600;color:{colour(value)};">'
            f'{value:.2f}</td>')


def matrix(title, regions, bases, cells, note=""):
    keys = [k for k in regions if any((b, k) in cells for b in bases)]
    head = "".join(f'<th style="{TH}text-align:center;">{regions[k]}</th>' for k in keys)
    body = []
    for i, base in enumerate(bases):
        if not any((base, k) in cells for k in keys):
            continue
        bg = f"background:{STRIPE};" if i % 2 else ""
        tds = "".join(pct_cell(cells.get((base, k))) for k in keys)
        body.append(f'<tr style="{bg}"><td style="{TD}font-weight:600;white-space:nowrap;">{base}</td>{tds}</tr>')
    return (f'<h3 style="color:{NAVY};margin:22px 0 6px;">{title}</h3>{note}'
            f'<table cellpadding="0" cellspacing="0" style="border-collapse:collapse;{FONT}">'
            f'<tr><th style="{TH}text-align:left;">Canary</th>{head}</tr>{"".join(body)}</table>')


def is_running(c):
    return c["state"] == "RUNNING" and c["runs"] and c["pct"] is not None


def attention_table(cells):
    """One row per running canary below AMBER: its regions, and each distinct top reason with the regions
    that hit it."""
    groups = defaultdict(list)
    for (base, region), c in cells.items():
        if is_running(c) and float(c["pct"]) < AMBER:
            groups[base].append((region, c))
    order = sorted(groups.items(), key=lambda kv: -sum(c["failed"] for _, c in kv[1]))
    body = []
    for i, (base, items) in enumerate(order):
        items.sort(key=lambda rc: list(ALL_REGIONS).index(rc[0]))
        regions = ", ".join(f'{ALL_REGIONS[r]} <span style="color:{colour(float(c["pct"]))};">'
                            f'({float(c["pct"]):.1f}%)</span>' for r, c in items)
        by_reason = defaultdict(list)
        for r, c in items:
            by_reason[c["reason"] or "no failure reason found in logs"].append(ALL_REGIONS[r])
        reasons = "".join(f'<div style="margin-bottom:4px;"><b style="font-family:Calibri,Arial,sans-serif;">'
                          f'{", ".join(rs)}:</b> {html.escape(reason)}</div>'
                          for reason, rs in by_reason.items())
        failed = sum(c["failed"] for _, c in items)
        bg = f"background:{STRIPE};" if i % 2 else ""
        body.append(f'<tr style="{bg}"><td style="{TD}font-weight:600;white-space:nowrap;">{base}</td>'
                    f'<td style="{TD}">{regions}</td><td style="{TD}text-align:right;">{failed:,}</td>'
                    f'<td style="{TD}font-family:Consolas,monospace;font-size:12px;">{reasons}</td></tr>')
    return (f'<table cellpadding="0" cellspacing="0" style="border-collapse:collapse;{FONT}">'
            f'<tr><th style="{TH}">Canary</th><th style="{TH}">Regions (success %)</th>'
            f'<th style="{TH}">Failed runs</th><th style="{TH}">Top failure reason (canary logs)</th></tr>'
            f'{"".join(body)}</table>')


def build(path):
    start, end, tz_name, cells = load_report(path)
    month = start.strftime("%B %Y") if start.month == end.month else \
        f"{start:%d %b} &ndash; {end:%d %b %Y}"
    running = [c for c in cells.values() if is_running(c)]
    runs, passed = sum(c["runs"] for c in running), sum(c["passed"] for c in running)
    green = sum(float(c["pct"]) >= GREEN for c in running)
    amber = sum(AMBER <= float(c["pct"]) < GREEN for c in running)
    red = len(running) - green - amber
    regions = len({r for (_, r) in cells})

    suite_bases = list(dict.fromkeys(b for (b, r) in cells if r in SUITE_REGIONS and b not in DATA_SYNC))
    legend = (f'<p style="color:#607d8b;font-size:12px;margin:6px 0 0;">Cell colour: '
              f'<b style="color:#2e7d32;">green &ge; {GREEN}%</b>, '
              f'<b style="color:#b7791f;">amber {AMBER:g}&ndash;{GREEN}%</b>, '
              f'<b style="color:#c00000;">red &lt; {AMBER:g}%</b>; '
              f'&ndash; = canary not deployed in that region; stopped = no runs in the period.</p>')

    return f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>Synthetic Canary Stability &mdash; {month}</title></head>
<body style="margin:0;padding:16px;{FONT}font-size:14px;color:#1f2937;line-height:1.5;background:#ffffff;">
<p>Hi All,</p>
<p>Please find below the CloudWatch Synthetics canary stability report for <b>{month}</b>
({start:%d %b} &ndash; {end:%d %b %Y}, {tz_name}) across {regions} production regions.</p>

<div style="background:#eef2f9;border-left:4px solid {NAVY};padding:10px 14px;margin:12px 0;">
<b style="color:{NAVY};">Summary</b><br>
Overall success ratio: <b>{passed / runs * 100:.2f}%</b> ({passed:,} of {runs:,} canary runs passed)<br>
Canaries running: <b>{len(running)}</b> &nbsp;&middot;&nbsp;
<b style="color:#2e7d32;">{green} green</b> &nbsp;&middot;&nbsp;
<b style="color:#b7791f;">{amber} amber</b> &nbsp;&middot;&nbsp;
<b style="color:#c00000;">{red} red</b>
</div>

{matrix("Success % by region &mdash; full canary suite", SUITE_REGIONS, suite_bases, cells)}
{matrix("Success % by region &mdash; data sync canaries", ALL_REGIONS, DATA_SYNC, cells)}
{legend}

<h3 style="color:{NAVY};margin:22px 0 6px;">Canaries below {AMBER:g}% &mdash; needs attention</h3>
{attention_table(cells)}

<p style="margin-top:16px;">The full per-region / per-canary breakdown, daily pass rates and every failure
reason are in the attached workbook.</p>
<p>Regards,</p>
</body></html>
"""


def main():
    args = parse_args()
    out = args.output or re.sub(r"\.xlsx$", "", args.report) + "_email.html"
    with open(out, "w", encoding="utf-8") as f:
        f.write(build(args.report))
    print(f"Email HTML written to {out}")


if __name__ == "__main__":
    main()
