"""
Fetch CloudWatch Synthetics canary pass rate and failure reasons for a date range,
across multiple CXone production regions, and write an Excel report.

Usage examples:
    python fetch_canary_report.py --start 2026-09-01 --end 2026-09-23
    python fetch_canary_report.py --start 2026-09-01 --end 2026-09-23 --regions na1 de
    python fetch_canary_report.py --start 2026-09-01 --end 2026-09-23 --regions jp jo ca kr na3 uae1 eu3 uk3
    python fetch_canary_report.py --start 2026-09-01 --end 2026-09-23 --canaries ua-curr-can ga-n-lg-can

Data sources:
    pass rate        CloudWatch metric CloudWatchSynthetics/SuccessPercent (kept 15 months)
    failure reasons  "Step failed" lines in the canary's Lambda logs, via Logs Insights
                     (log retention is typically 90 days); falls back to GetCanaryRuns
                     history, which only holds the last few days.

Dates are inclusive and interpreted in UTC unless --tz is given (e.g. --tz IST for UTC+05:30).
"""

import argparse
import csv
import glob
import json
import os
import re
import sys
import time
from collections import defaultdict
from datetime import datetime, timedelta, timezone

import boto3
from botocore.exceptions import BotoCoreError, ClientError
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CANARIES = [
    "gla-canary", "ga-o-lg-can", "ga-o-pm-can", "ga-oldal-can", "ua-old-can",
    "ua-light-can", "ga-ds-red-can", "ga-cural-can", "ga-dspwd-can", "ua-curr-can",
    "ga-c-pm-can", "ga-light-can", "ga-c-lg-can", "ua-new-can", "ga-n-pm-can",
    "ga-newal-can", "ga-n-lg-can",
]

DS_CANARIES = ["ga-dsync-can", "ga-dspwd-can"]

UK3_CANARIES = [
    "ua-light-can", "ua-curr-can", "ua-old-can", "ga-dsync-can", "ga-c-pm-can", "ga-o-pm-can",
    "ga-n-pm-can", "gla-canary", "ga-ds-red-can", "ga-n-lg-can", "ga-c-lg-can", "ga-o-lg-can",
    "ga-new-can", "ga-newal-can", "ga-light-can",
]

# region key -> AWS region, canary name prefixes (first existing one wins), owning AWS account,
# optional canary list (default CANARIES) and report table (each table is its own success % matrix).
# Credentials are resolved per region: --profile override if given, otherwise the default
# credentials if they belong to the right account, otherwise the SSO profile wfo-acct-<account>,
# otherwise the cached `aws sso login` token (see sso_session).
REGIONS = {
    "na1":   {"aws_region": "us-west-2",      "prefixes": ["prod-"],                       "account": "918987959928"},
    "de":    {"aws_region": "eu-central-1",   "prefixes": ["prod-de-"],                    "account": "918987959928"},
    "au":    {"aws_region": "ap-southeast-2", "prefixes": ["prod-au-"],                    "account": "918987959928"},
    "ausov": {"aws_region": "ap-southeast-2", "prefixes": ["au2-", "sau2-", "prod-au2-"], "account": "637423329662"},
    "eusov": {"aws_region": "eu-central-1",   "prefixes": ["eu2-", "seu2-", "prod-eu2-"], "account": "851725389895"},
    "uksov": {"aws_region": "eu-west-2",      "prefixes": ["uk2-", "suk2-", "prod-uk2-"], "account": "571600874006"},
    "fedramp": {"aws_region": "us-east-1",    "prefixes": ["fedr-", "prod-fed-"],          "account": "334442430111"},
    "uk":    {"aws_region": "eu-west-2",      "prefixes": ["prod-uk-"],                    "account": "918987959928",
              "canaries": DS_CANARIES, "table": "other"},
    "jp":    {"aws_region": "ap-northeast-1", "prefixes": ["prod-jp-"],                    "account": "918987959928",
              "canaries": DS_CANARIES, "table": "other"},
    "jo":    {"aws_region": "ap-northeast-3", "prefixes": ["prod-jo-"],                    "account": "692416605838",
              "canaries": DS_CANARIES, "table": "other"},
    "ca":    {"aws_region": "ca-central-1",   "prefixes": ["prod-ca-"],                    "account": "918987959928",
              "canaries": DS_CANARIES, "table": "other"},
    "kr":    {"aws_region": "ap-northeast-2", "prefixes": ["prod-kr-"],                    "account": "814740517824",
              "canaries": DS_CANARIES, "table": "other"},
    "na3":   {"aws_region": "us-east-1",      "prefixes": ["na3-"],                        "account": "686255955190",
              "canaries": DS_CANARIES, "table": "other"},
    "uae1":  {"aws_region": "me-central-1",   "prefixes": ["prod-ae-"],                    "account": "371021933542",
              "canaries": DS_CANARIES, "table": "other"},
    "za":    {"aws_region": "af-south-1",     "prefixes": ["za1-", "prod-za-"],            "account": "529088275443",
              "canaries": DS_CANARIES, "table": "other"},
    "eu3":   {"aws_region": "eu-west-1",      "prefixes": ["eu3-"],                        "account": "557690609302",
              "canaries": DS_CANARIES, "table": "other"},
    "uk3":   {"aws_region": "eu-west-2",      "prefixes": ["uk3-", "suk3-", "prod-uk3-"], "account": "842477702373",
              "canaries": UK3_CANARIES, "table": "other"},
}
TABLES = {"main": "Success %", "other": "Success % (other regions)"}
# Fallback when no profile works: credentials straight from the cached `aws sso login` token.
SSO_START_URL = "https://identitycenter.amazonaws.com/ssoins-72233ac939008b91"
SSO_ROLE = "Dept_Dev-Platform-SnC"

DATE_FMT = "%Y-%m-%d"
LOGS_QUERY = (
    "fields @timestamp, @message "
    "| filter @message like /Step failed: / or @message like /status: FAILED failureReason:/ "
    "or @message like /Task timed out after/ "
    "| sort @timestamp asc | limit 10000"
)
MAX_CONCURRENT_QUERIES = 10


TZ_ALIASES = {"UTC": "+00:00", "IST": "+05:30"}
# Report time zone, set from --tz in main(); used for dates and every timestamp written to the report.
TZ, TZ_NAME = timezone.utc, "UTC"


def parse_tz(value):
    """'UTC', 'IST' or '+HH:MM' / '-HH:MM' -> (tzinfo, display name)."""
    offset = TZ_ALIASES.get(value.upper(), value)
    m = re.fullmatch(r"([+-])(\d{1,2}):?(\d{2})", offset)
    if not m:
        raise argparse.ArgumentTypeError(f"bad time zone {value!r} (use UTC, IST or +HH:MM)")
    delta = timedelta(hours=int(m.group(2)), minutes=int(m.group(3))) * (-1 if m.group(1) == "-" else 1)
    name = value.upper() if value.upper() in TZ_ALIASES else f"UTC{offset}"
    if name == "IST":
        name = "IST, UTC+05:30"
    return timezone(delta), name


def local(ts):
    """datetime -> 'YYYY-MM-DD HH:MM:SS' in the report time zone."""
    return ts.astimezone(TZ).strftime("%Y-%m-%d %H:%M:%S")


def parse_args():
    p = argparse.ArgumentParser(description="Canary pass-rate / failure-reason report")
    p.add_argument("--start", required=True, help="Start date YYYY-MM-DD (inclusive, in --tz)")
    p.add_argument("--end", required=True, help="End date YYYY-MM-DD (inclusive, in --tz)")
    p.add_argument("--tz", default="UTC", type=parse_tz,
                   help="Time zone for the date range, daily buckets and timestamps: UTC (default), IST, "
                        "or an offset like +05:30")
    p.add_argument("--regions", nargs="+", default=list(REGIONS), choices=list(REGIONS),
                   help="Region keys to include (default: all)")
    p.add_argument("--canaries", nargs="+",
                   help="Base canary names (default: each region's own list; CANARIES unless the region sets one)")
    p.add_argument("--profile", action="append", default=[], metavar="REGION=PROFILE",
                   help="Override AWS profile for a region, e.g. de=production-de (repeatable)")
    p.add_argument("--output", help="Output .xlsx path (default: canary_report_<start>_<end>.xlsx)")
    return p.parse_args()


UUID_RE = re.compile(r"\b[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\b")


def clean(text):
    text = UUID_RE.sub("<id>", text)
    text = re.sub(r"\d{4}-\d{2}-\d{2}T[\d:.]+Z?", "<time>", text)
    return text.strip(" |\"")


def reason_from_log(message):
    """Turn a failed-step / Lambda timeout log line into a groupable reason.

    Older Synthetics runtimes log 'Step failed: <step> URL: <url> <error>'; newer ones log JSON with
    'Finished executing Step N: <step> status: FAILED failureReason: '<error>''.
    """
    m = re.search(r"Task timed out after [\d.]+ seconds", message)
    if m:
        return "Lambda timeout: " + re.sub(r"[\d.]+ seconds", "N seconds", m.group(0))
    if "failureReason:" in message:
        text = message
        brace = message.find("{")
        if brace >= 0:
            try:
                text = str(json.loads(message[brace:]).get("message", message))
            except ValueError:
                pass
        m = re.search(r"Step \d+: (?P<step>.*?)\s+status: FAILED failureReason: (?P<err>.*)", text, re.S)
        if m:
            err = m.group("err").replace("\\n", "\n").splitlines()[0].split(" Stack:")[0].strip(" '")
            return clean(f"[{m.group('step').strip()}] {err}")[:400]
    body = message.split("Step failed: ", 1)[-1].replace("\n", " ")
    m = re.match(r"(?P<step>.*?) URL: (?P<url>\S+)\s*(?P<err>.*)", body)
    if not m:
        return clean(body.split(" Stack:")[0])[:400]
    err = m.group("err").split(" Stack:")[0].split("    at ")[0]
    url = m.group("url").split("?")[0]
    return clean(f"[{m.group('step')}] {err} (at {url})")[:400]


def reason_from_run(reason):
    if not reason:
        return "(no reason given)"
    return clean(reason.strip().splitlines()[0].split(" Stack:")[0].split("\\n")[0])[:400]


def fetch_runs(syn, name, start, end):
    """Return runs whose start time is within [start, end). Runs come back newest first."""
    runs, token = [], None
    while True:
        kwargs = {"Name": name, "MaxResults": 100}
        if token:
            kwargs["NextToken"] = token
        resp = syn.get_canary_runs(**kwargs)
        page = resp.get("CanaryRuns", [])
        for r in page:
            started = r.get("Timeline", {}).get("Started")
            if started and start <= started < end:
                runs.append(r)
        token = resp.get("NextToken")
        oldest = min((r["Timeline"]["Started"] for r in page if r.get("Timeline", {}).get("Started")), default=None)
        if not token or (oldest and oldest < start):
            return runs


def fetch_metric(cw, name, start, end):
    """Daily (runs, passed) from the SuccessPercent metric."""
    resp = cw.get_metric_statistics(
        Namespace="CloudWatchSynthetics", MetricName="SuccessPercent",
        Dimensions=[{"Name": "CanaryName", "Value": name}],
        StartTime=start, EndTime=end, Period=86400, Statistics=["Average", "SampleCount"],
    )
    daily = {}
    for dp in resp.get("Datapoints", []):
        n = int(dp["SampleCount"])
        # Buckets start at StartTime (local midnight), so the local date of the bucket start is the day.
        daily[dp["Timestamp"].astimezone(TZ).strftime(DATE_FMT)] = (n, round(dp["Average"] / 100 * n))
    return daily


def log_groups_for(logs, name):
    """All non-empty Lambda log groups for a canary (one per engine the canary has used)."""
    groups, token = [], None
    while True:
        kwargs = {"logGroupNamePrefix": f"/aws/lambda/cwsyn-{name}-"}
        if token:
            kwargs["nextToken"] = token
        resp = logs.describe_log_groups(**kwargs)
        groups += [g["logGroupName"] for g in resp["logGroups"] if g.get("storedBytes", 1) > 0]
        token = resp.get("nextToken")
        if not token:
            return groups


def fetch_log_failures(logs, jobs, start, end):
    """Run Logs Insights for each (name, log_groups) job; return {name: [(timestamp, run_id, reason)] or error str}."""
    results, pending, queue = {}, {}, list(jobs)
    while queue or pending:
        while queue and len(pending) < MAX_CONCURRENT_QUERIES:
            name, groups = queue.pop(0)
            try:
                qid = logs.start_query(logGroupNames=groups, startTime=int(start.timestamp()),
                                       endTime=int(end.timestamp()), queryString=LOGS_QUERY)["queryId"]
                pending[qid] = name
            except ClientError as e:
                if e.response["Error"]["Code"] == "LimitExceededException":
                    queue.insert(0, (name, groups))
                    break
                results[name] = f"logs query failed: {e}"
        time.sleep(3)
        for qid, name in list(pending.items()):
            resp = logs.get_query_results(queryId=qid)
            if resp["status"] in ("Scheduled", "Running"):
                continue
            del pending[qid]
            if resp["status"] != "Complete":
                results[name] = f"logs query {resp['status']}"
                continue
            rows = []
            for r in resp["results"]:
                f = {x["field"]: x["value"] for x in r}
                run = UUID_RE.search(f["@message"])  # canaryRunId (JSON logs) or Lambda request id
                logged = datetime.strptime(f["@timestamp"][:19], "%Y-%m-%d %H:%M:%S").replace(tzinfo=timezone.utc)
                rows.append((local(logged), run.group(0) if run else f["@timestamp"],
                             reason_from_log(f["@message"])))
            results[name] = rows
    return results


def cached_sso_token():
    """Return (access token, SSO region) of the newest unexpired token from `aws sso login`, or None.

    Any SSO session or profile login counts, so one login covers every account the user can reach
    even when a profile's own sso_session/sso_start_url points at an expired login.
    """
    best = None
    now = datetime.now(timezone.utc)
    for path in glob.glob(os.path.expanduser(os.path.join("~", ".aws", "sso", "cache", "*.json"))):
        try:
            with open(path, encoding="utf-8") as f:
                data = json.load(f)
            expires = datetime.fromisoformat(data["expiresAt"].replace("Z", "+00:00"))
            token, region = data["accessToken"], data["region"]
        except (OSError, ValueError, KeyError):
            continue
        if data.get("startUrl") == SSO_START_URL and expires > now and (not best or expires > best[0]):
            best = (expires, token, region)
    return best[1:] if best else None


def sso_session(cfg):
    """Session for cfg["account"] built directly from the cached SSO token, or raise RuntimeError."""
    cached = cached_sso_token()
    if not cached:
        raise RuntimeError("no unexpired SSO token in ~/.aws/sso/cache (run aws sso login)")
    token, sso_region = cached
    sso = boto3.client("sso", region_name=sso_region)
    try:
        roles = [r["roleName"] for r in sso.list_account_roles(
            accessToken=token, accountId=cfg["account"])["roleList"]]
        if not roles:
            raise RuntimeError(f"SSO user has no role in account {cfg['account']}")
        role = SSO_ROLE if SSO_ROLE in roles else roles[0]
        creds = sso.get_role_credentials(roleName=role, accountId=cfg["account"],
                                         accessToken=token)["roleCredentials"]
    except (ClientError, BotoCoreError) as e:
        raise RuntimeError(f"SSO token: {e}") from e
    session = boto3.Session(aws_access_key_id=creds["accessKeyId"],
                            aws_secret_access_key=creds["secretAccessKey"],
                            aws_session_token=creds["sessionToken"], region_name=cfg["aws_region"])
    return session, f"sso-token:{role}"


def open_session(cfg, override):
    """Return (session, profile, account) for credentials that belong to cfg["account"]."""
    candidates = [override] if override else [None, f"wfo-acct-{cfg['account']}"]
    problems = []
    for profile in candidates:
        try:
            session = boto3.Session(profile_name=profile, region_name=cfg["aws_region"])
            account = session.client("sts").get_caller_identity()["Account"]
        except (ClientError, BotoCoreError) as e:
            problems.append(f"{profile or 'default'}: {e}")
            continue
        if account == cfg["account"]:
            return session, profile, account
        problems.append(f"{profile or 'default'}: credentials are for account {account}")
    if not override:
        try:
            session, profile = sso_session(cfg)
            return session, profile, cfg["account"]
        except RuntimeError as e:
            problems.append(str(e))
    raise RuntimeError(f"no credentials for account {cfg['account']} ({'; '.join(problems)})")


def list_canaries(syn):
    existing, token = {}, None
    while True:
        resp = syn.describe_canaries(**({"NextToken": token} if token else {}))
        existing.update((c["Name"], c) for c in resp["Canaries"])
        token = resp.get("NextToken")
        if not token:
            return existing


def region_canaries(cfg, args):
    """Base canary names to report for a region; --canaries narrows a region with its own list."""
    if "canaries" not in cfg:
        return args.canaries or CANARIES
    return [b for b in cfg["canaries"] if not args.canaries or b in args.canaries]


def pct(passed, total):
    return round(passed / total * 100, 2) if total else None


def blank_row(key, cfg, account, name, state, note):
    return [key, cfg["aws_region"], account, name, state, "", "", "", "", "", "", "", note]


def process_region(key, cfg, args, overrides, start, end, out):
    print(f"\n=== {key} ({cfg['aws_region']}, account {cfg['account']}) ===")
    try:
        session, profile, account = open_session(cfg, overrides.get(key))
        syn, cw, logs = (session.client(s) for s in ("synthetics", "cloudwatch", "logs"))
        existing = list_canaries(syn)
    except (RuntimeError, ClientError, BotoCoreError) as e:
        print(f"  SKIPPED: {e}")
        for base in region_canaries(cfg, args):
            out["summary"].append(blank_row(key, cfg, cfg["account"], cfg["prefixes"][0] + base,
                                            "NO ACCESS", str(e)[:300]))
            out["matrix"][(base, key)] = "NO ACCESS"
        return
    print(f"  using profile {profile or 'default'}, {len(existing)} canaries in region")

    found = []
    for base in region_canaries(cfg, args):
        name = next((p + base for p in cfg["prefixes"] if p + base in existing), None)
        if name:
            found.append((base, name))
            continue
        similar = ", ".join(sorted(n for n in existing if n.endswith(base))) or "-"
        print(f"  {cfg['prefixes'][0] + base}: NOT FOUND (similar: {similar})")
        out["summary"].append(blank_row(key, cfg, account, cfg["prefixes"][0] + base, "NOT FOUND",
                                        f"Similar: {similar}"))
        out["matrix"][(base, key)] = "N/A"

    jobs, log_notes = [], {}
    for _, name in found:
        try:
            groups = log_groups_for(logs, name)
        except (ClientError, BotoCoreError) as e:
            groups, log_notes[name] = [], f"cannot list log groups: {e}"
        if groups:
            jobs.append((name, groups[:50]))
        else:
            log_notes.setdefault(name, "no Lambda log group")
    print(f"  querying logs for {len(jobs)} canaries...")
    log_results = fetch_log_failures(logs, jobs, start, end)

    for base, name in found:
        canary = existing[name]
        try:
            runs = fetch_runs(syn, name, start, end)
            daily = fetch_metric(cw, name, start, end)
        except (ClientError, BotoCoreError) as e:
            print(f"  {name}: ERROR {e}")
            out["summary"].append(blank_row(key, cfg, account, name, canary["Status"]["State"], str(e)[:300]))
            out["matrix"][(base, key)] = "ERROR"
            continue

        history = [r for r in runs if r["Status"]["State"] in ("PASSED", "FAILED")]
        total = sum(n for n, _ in daily.values())
        passed = sum(p for _, p in daily.values())
        source = "metric"
        if not total and history:
            total = len(history)
            passed = sum(1 for r in history if r["Status"]["State"] == "PASSED")
            source = "run history"

        failures = log_results.get(name)
        notes = [f"pass rate from {source}"]
        if isinstance(failures, list):
            reason_source = "logs"
        else:
            notes.append(failures or log_notes.get(name, "no logs"))
            failures = [(local(r["Timeline"]["Started"]), r["Id"], reason_from_run(r["Status"].get("StateReason")))
                        for r in history if r["Status"]["State"] == "FAILED"]
            reason_source = "run history (recent days only)"

        # A run can fail several steps, so count each reason once per run.
        grouped = defaultdict(dict)
        for ts, run, reason in failures:
            grouped[reason].setdefault(run, ts)
        failed_runs = len({run for _, run, _ in failures})
        ranked = sorted(grouped.items(), key=lambda kv: -len(kv[1]))
        for reason, runs_ts in ranked:
            stamps = runs_ts.values()
            out["reasons"].append([key, name, len(runs_ts), pct(len(runs_ts), failed_runs),
                                   min(stamps), max(stamps), reason_source, reason])
        for ts, run, reason in failures:
            out["failures"].append([key, name, ts, run, reason_source, reason])
        for day, (n, p) in sorted(daily.items()):
            out["daily"].append([key, name, day, n, p, n - p, pct(p, n)])
        for r in runs:
            out["runs"].append([key, name, local(r["Timeline"]["Started"]), r["Status"]["State"], r["Status"].get("StateReason", "")[:1000]])

        out["matrix"][(base, key)] = pct(passed, total) if total else "NO RUNS"
        top = f"{ranked[0][0]} ({len(ranked[0][1])}x)" if ranked else ""
        out["summary"].append([key, cfg["aws_region"], account, name, canary["Status"]["State"],
                               canary["Schedule"]["Expression"], total, passed, total - passed,
                               pct(passed, total), failed_runs, top, "; ".join(notes)])
        print(f"  {name}: runs={total} passed={passed} failed={total - passed} "
              f"pass%={pct(passed, total)} | failed runs with reason ({reason_source}): {failed_runs}")


def main():
    global TZ, TZ_NAME
    args = parse_args()
    TZ, TZ_NAME = args.tz
    start = datetime.strptime(args.start, DATE_FMT).replace(tzinfo=TZ)
    end = datetime.strptime(args.end, DATE_FMT).replace(tzinfo=TZ) + timedelta(days=1)
    print(f"Range: {start:%Y-%m-%d %H:%M} to {end:%Y-%m-%d %H:%M} {TZ_NAME} "
          f"(UTC {start.astimezone(timezone.utc):%Y-%m-%d %H:%M} to {end.astimezone(timezone.utc):%Y-%m-%d %H:%M})")
    overrides = dict(o.split("=", 1) for o in args.profile)
    out = {k: [] for k in ("summary", "reasons", "failures", "daily", "runs")}
    out["matrix"] = {}

    for key in args.regions:
        process_region(key, REGIONS[key], args, overrides, start, end, out)

    suffix = "" if TZ_NAME == "UTC" else "_" + TZ_NAME.split(",")[0].replace("+", "p").replace(":", "")
    path = args.output or f"canary_report_{args.start}_{args.end}{suffix}.xlsx"
    write_workbook(path, args, out)
    print(f"\nReport written to {path}")
    matrix_path = re.sub(r"\.xlsx$", "", path) + "_success_pct.csv"
    write_matrix_csv(matrix_path, args, out["matrix"])
    print(f"Success % matrix written to {matrix_path}")


def matrix_tables(args, matrix):
    """One (title, header, rows) success % matrix per table: canary (rows) x region (columns).

    N/A = canary not found in region, NO RUNS = no runs in range, blank = canary not listed for region.
    """
    tables = []
    for table, title in TABLES.items():
        keys = [k for k in args.regions if REGIONS[k].get("table", "main") == table]
        if not keys:
            continue
        bases = list(dict.fromkeys(b for k in keys for b in region_canaries(REGIONS[k], args)))
        rows = [[base] + [matrix.get((base, key), "") for key in keys] for base in bases]
        tables.append((title, ["Canary"] + keys, rows))
    return tables


def write_matrix_csv(path, args, matrix):
    """All success % matrices in one CSV, separated by a blank line and a title row."""
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for i, (title, header, rows) in enumerate(matrix_tables(args, matrix)):
            if i:
                w.writerows([[], [title]])
            w.writerow(header)
            w.writerows(rows)


def write_sheet(ws, headers, rows, widths):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str) and len(cell.value) > 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_workbook(path, args, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    write_sheet(ws, ["Region", "AWS Region", "Account", "Canary", "State", "Schedule",
                     "Runs", "Passed", "Failed", "Pass %", "Failed runs with reason",
                     "Top failure reason", "Notes"],
                out["summary"], [8, 15, 14, 24, 11, 22, 8, 8, 8, 8, 10, 90, 40])
    for title, header, rows in matrix_tables(args, out["matrix"]):
        write_sheet(wb.create_sheet(title), header, rows, [18] + [10] * (len(header) - 1))
    write_sheet(wb.create_sheet("Failure Reasons"),
                ["Region", "Canary", "Failed runs", "% of failed runs", f"First seen ({TZ_NAME})",
                 f"Last seen ({TZ_NAME})", "Source", "Reason"],
                out["reasons"], [8, 24, 8, 10, 19, 19, 14, 120])
    write_sheet(wb.create_sheet("Daily"),
                ["Region", "Canary", f"Date ({TZ_NAME})", "Runs", "Passed", "Failed", "Pass %"],
                out["daily"], [8, 24, 12, 8, 8, 8, 8])
    write_sheet(wb.create_sheet("Failures"),
                ["Region", "Canary", f"Time ({TZ_NAME})", "Run / request id", "Source", "Reason"],
                out["failures"], [8, 24, 20, 38, 14, 120])
    write_sheet(wb.create_sheet("Recent Runs"),
                ["Region", "Canary", f"Started ({TZ_NAME})", "Status", "Reason"],
                out["runs"], [8, 24, 20, 10, 120])
    info = wb.create_sheet("Info")
    for line in [
        f"Date range ({TZ_NAME}, inclusive): {args.start} to {args.end}",
        f"Generated: {datetime.now(timezone.utc):%Y-%m-%d %H:%M:%S} UTC",
        "Runs / Passed / Pass % come from the CloudWatch SuccessPercent metric (one sample per run).",
        "Failure reasons come from 'Step failed' and Lambda timeout lines in the canary Lambda logs.",
        "A run can fail several steps; 'Failure Reasons' counts each reason once per run, so percentages can sum past 100%.",
        "'Recent Runs' is GetCanaryRuns history, which only keeps the last few days (canary retention setting).",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 110
    wb.save(path)


if __name__ == "__main__":
    sys.exit(main())
